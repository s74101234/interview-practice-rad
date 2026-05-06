import pymupdf4llm
from pptx import Presentation
from pathlib import Path


def parse_pdf(file_path: str) -> list[dict]:
    pages = pymupdf4llm.to_markdown(file_path, page_chunks=True)
    return [
        {"source": Path(file_path).name, "page": str(p["metadata"]["page"] + 1), "content": p["text"]}
        for p in pages
        if p["text"].strip()
    ]


def parse_pptx(file_path: str) -> list[dict]:
    prs = Presentation(file_path)
    results = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        texts.append(text)
        if texts:
            results.append({
                "source": Path(file_path).name,
                "page": str(i),
                "content": "\n".join(texts),
            })
    return results


def parse(file_path: str) -> list[dict]:
    ext = Path(file_path).suffix.lower()
    if ext == ".pdf":
        return parse_pdf(file_path)
    elif ext in (".pptx", ".ppt"):
        return parse_pptx(file_path)
    raise ValueError(f"Unsupported file type: {ext}")
